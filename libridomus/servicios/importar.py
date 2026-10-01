"""Importar elementos desde una hoja de cálculo (Excel .xlsx) o un archivo CSV.

Pasos:
1. ``leer_tabla`` lee el archivo: la primera fila son los nombres de las columnas.
2. ``sugerir_destino`` propone a qué dato de LibriDomus va cada columna (por su nombre).
3. ``importar`` crea los elementos dentro de una sola transacción y devuelve un informe
   con lo creado, lo omitido y los avisos (fila por fila, para que el usuario pueda corregir).

Nada se pierde en silencio: lo que no encaja en ningún campo (una conservación desconocida,
un campo que el tipo no tiene, una fecha mal escrita...) se añade a las notas del elemento.
"""

import csv
import io
import re
import sqlite3
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from .. import texto
from ..datos import elementos, tipos, ubicaciones
from ..datos.conexion import transaccion
from ..datos.elementos import Elemento, ErrorElemento

MAX_BYTES = 50 * 1024 * 1024    # archivos más grandes no son una colección doméstica: se rechazan
MAX_FILAS = 200_000
MAX_CELDA = 10_000              # caracteres por celda (lo demás se recorta)

IGNORAR = ""
PREFIJO_CAMPO = "campo:"        # destinos de campos propios: 'campo:<etiqueta>'

# Datos comunes a todos los tipos: (clave, nombre que ve el usuario)
DESTINOS_COMUNES = [
    ("titulo", "Título"),
    ("subtitulo", "Subtítulo"),
    ("personas", "Personas (autor, director…)"),
    ("anio", "Año"),
    ("identificador", "ISBN / EAN / identificador"),
    ("tipo", "Tipo de elemento"),
    ("ubicacion", "Ubicación (código o ruta)"),
    ("idioma", "Idioma"),
    ("estado", "Conservación"),
    ("valoracion", "Valoración (0 a 5)"),
    ("consumido", "Leído / visto / escuchado"),
    ("etiquetas", "Etiquetas"),
    ("prestado_a", "Prestado a"),
    ("fecha_desde", "Periodo desde"),
    ("fecha_hasta", "Periodo hasta"),
    ("lugar_evento", "Lugar / evento"),
    ("notas", "Notas"),
]

# Nombres de columna habituales (sin acentos ni mayúsculas) -> destino
SINONIMOS = {
    "titulo": "titulo", "title": "titulo", "nombre": "titulo", "obra": "titulo",
    "subtitulo": "subtitulo", "subtitle": "subtitulo",
    "autor": "personas", "autores": "personas", "author": "personas", "authors": "personas",
    "autor/a": "personas", "personas": "personas", "artista": "personas", "interprete": "personas",
    "director": "personas", "compositor": "personas", "guionista": "personas", "dibujante": "personas",
    "traductor": "personas", "ilustrador": "personas",
    "ano": "anio", "anio": "anio", "year": "anio", "fecha": "anio",
    "ano de publicacion": "anio", "fecha de publicacion": "anio",
    "isbn": "identificador", "ean": "identificador", "issn": "identificador", "identificador": "identificador",
    "codigo de barras": "identificador", "isbn13": "identificador", "isbn-13": "identificador",
    "tipo": "tipo", "clase": "tipo", "categoria": "tipo",
    "ubicacion": "ubicacion", "lugar": "ubicacion", "estanteria": "ubicacion", "codigo": "ubicacion",
    "donde": "ubicacion", "sitio": "ubicacion", "balda": "ubicacion", "caja": "ubicacion",
    "idioma": "idioma", "lengua": "idioma", "language": "idioma",
    "estado": "estado", "conservacion": "estado",
    "valoracion": "valoracion", "puntuacion": "valoracion", "nota": "valoracion", "estrellas": "valoracion",
    "leido": "consumido", "visto": "consumido", "escuchado": "consumido", "jugado": "consumido",
    "etiquetas": "etiquetas", "etiqueta": "etiquetas", "tags": "etiquetas", "genero": "etiquetas",
    "prestado a": "prestado_a", "prestado": "prestado_a", "prestamo": "prestado_a",
    "desde": "fecha_desde", "periodo desde": "fecha_desde", "hasta": "fecha_hasta",
    "periodo hasta": "fecha_hasta", "evento": "lugar_evento", "lugar / evento": "lugar_evento",
    "notas": "notas", "observaciones": "notas", "comentarios": "notas", "comentario": "notas",
}

VALORES_SI = {"si", "sí", "s", "x", "1", "true", "verdadero", "yes", "y", "leido", "visto", "escuchado", "jugado"}


class ErrorImportar(Exception):
    """Archivo que no se puede leer (mensaje para el usuario)."""


@dataclass
class Tabla:
    cabeceras: list[str]
    filas: list[list[str]]


@dataclass
class Informe:
    creados: list[int] = field(default_factory=list)
    omitidos: list[tuple[int, str]] = field(default_factory=list)   # (nº de fila en el archivo, motivo)
    avisos: list[tuple[int, str]] = field(default_factory=list)


def normalizar(valor: str) -> str:
    """'  Año de Publicación ' -> 'ano de publicacion' (sin acentos, minúsculas, espacios simples)."""
    return re.sub(r"\s+", " ", texto.sin_acentos(str(valor or "")).casefold()).strip()


# ------------------------------------------------------------------ lectura del archivo

def _celda(valor) -> str:
    """Convierte lo que da Excel en texto: 2004.0 -> '2004', fechas -> 'AAAA-MM-DD'."""
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "Sí" if valor else ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    if isinstance(valor, datetime):
        valor = valor.date()
    if isinstance(valor, date):
        return valor.isoformat()
    return elementos.limpiar_texto(valor).strip()[:MAX_CELDA]


def _leer_csv(ruta: Path) -> list[list[str]]:
    datos = ruta.read_bytes()
    for codificacion in ("utf-8-sig", "cp1252"):  # Excel en español guarda los CSV en ANSI (cp1252)
        try:
            contenido = datos.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    else:
        contenido = datos.decode("utf-8", "replace")
    muestra = contenido[:20_000]
    try:
        dialecto = csv.Sniffer().sniff(muestra, delimiters=";,\t|")
        separador = dialecto.delimiter
    except csv.Error:
        separador = ";" if muestra.count(";") >= muestra.count(",") else ","
    lector = csv.reader(io.StringIO(contenido, newline=""), delimiter=separador)
    filas = []
    for fila in lector:
        filas.append([_celda(c) for c in fila])
        if len(filas) > MAX_FILAS + 1:
            raise ErrorImportar(f"El archivo tiene más de {MAX_FILAS:,} filas.".replace(",", "."))
    return filas


def _leer_xlsx(ruta: Path) -> list[list[str]]:
    import openpyxl  # solo se carga si hace falta (tarda un poco en importarse)
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    except (InvalidFileException, KeyError, ValueError, OSError) as e:
        raise ErrorImportar(f"No es un archivo de Excel válido ({e}).") from e
    except Exception as e:  # un .xlsx dañado puede fallar de muchas formas (zip, XML...)
        raise ErrorImportar(f"No se puede leer el archivo de Excel ({type(e).__name__}).") from e
    try:
        hoja = libro.worksheets[0]  # siempre la primera hoja
        filas = []
        for fila in hoja.iter_rows(values_only=True):
            filas.append([_celda(c) for c in fila])
            if len(filas) > MAX_FILAS + 1:
                raise ErrorImportar(f"La hoja tiene más de {MAX_FILAS:,} filas.".replace(",", "."))
        return filas
    finally:
        libro.close()


def leer_tabla(ruta: Path | str) -> Tabla:
    """Lee un .csv, .txt o .xlsx. La primera fila con algo escrito son las cabeceras."""
    ruta = Path(ruta)
    try:
        tamano = ruta.stat().st_size
    except OSError as e:
        raise ErrorImportar(f"No se puede abrir el archivo: {e}") from e
    if tamano > MAX_BYTES:
        raise ErrorImportar("El archivo es demasiado grande (más de 50 MB).")
    extension = ruta.suffix.lower()
    if extension in (".xlsx", ".xlsm"):
        filas = _leer_xlsx(ruta)
    elif extension in (".csv", ".txt"):
        try:
            filas = _leer_csv(ruta)
        except OSError as e:
            raise ErrorImportar(f"No se puede leer el archivo: {e}") from e
    elif extension == ".xls":
        raise ErrorImportar("Los archivos .xls (Excel 97-2003) no se pueden leer. Ábrelo en Excel y "
                            "guárdalo como «Libro de Excel (.xlsx)» o como CSV.")
    else:
        raise ErrorImportar("Elige un archivo de Excel (.xlsx) o CSV (.csv).")
    filas = [f for f in filas if any(c for c in f)]  # sin filas vacías
    if not filas:
        raise ErrorImportar("El archivo está vacío.")
    ancho = max(len(f) for f in filas)
    filas = [f + [""] * (ancho - len(f)) for f in filas]
    # Columnas sin cabecera ni datos (habituales en Excel) se quitan.
    usadas = [i for i in range(ancho) if any(f[i] for f in filas)]
    filas = [[f[i] for i in usadas] for f in filas]
    cabeceras = [c or f"Columna {n + 1}" for n, c in enumerate(filas[0])]
    return Tabla(cabeceras=cabeceras, filas=filas[1:])


# ------------------------------------------------------------------ correspondencia de columnas

def destinos(con: sqlite3.Connection) -> list[tuple[str, str]]:
    """Todos los destinos posibles: los comunes y los campos propios de todos los tipos."""
    resultado = list(DESTINOS_COMUNES)
    vistos = set()
    for t in tipos.listar(con):
        for c in t.campos:
            clave = normalizar(c.etiqueta)
            if clave not in vistos:
                vistos.add(clave)
                resultado.append((PREFIJO_CAMPO + clave, c.etiqueta))
    return resultado


def sugerir_destino(cabecera: str, disponibles: list[tuple[str, str]]) -> str:
    """Propone el destino de una columna según su nombre ('' si no se reconoce)."""
    nombre = normalizar(cabecera)
    if nombre in SINONIMOS:
        return SINONIMOS[nombre]
    for clave, etiqueta in disponibles:
        if nombre in (normalizar(etiqueta), normalizar(clave.removeprefix(PREFIJO_CAMPO))):
            return clave
    return IGNORAR


# ------------------------------------------------------------------ conversión de valores

def _anio(valor: str) -> int | None:
    m = re.search(r"\b(\d{1,4})\b", valor)
    return int(m.group(1)) if m else None


def _valoracion(valor: str) -> int | None:
    if "★" in valor:
        return min(5, valor.count("★"))
    m = re.fullmatch(r"\s*(\d)([.,]\d+)?\s*(/\s*5)?\s*", valor)
    if m and 0 <= int(m.group(1)) <= 5:
        return int(m.group(1))
    return None


def _personas(valor: str, rol_columna: str) -> list[tuple[str, str]]:
    """'Ana Pérez; Luis Gil (Traductor)' -> [('Ana Pérez', rol_columna), ('Luis Gil', 'Traductor')]."""
    resultado = []
    for parte in re.split(r"[;\n]", valor):
        parte = parte.strip()
        m = re.fullmatch(r"(.+?)\s*\(([^()]+)\)", parte)
        if m:
            resultado.append((m.group(1).strip(), m.group(2).strip()))
        elif parte:
            resultado.append((parte, rol_columna))
    return resultado


class _Ubicaciones:
    """Busca ubicaciones por código, por ruta ('Planta baja > Salón > Estantería A') o por nombre."""

    def __init__(self, con: sqlite3.Connection):
        self.por_codigo, self.por_ruta, nombres = {}, {}, {}
        rutas = ubicaciones.rutas_todas(con)
        for u in ubicaciones.todas(con):
            self.por_codigo[normalizar(u.codigo)] = u.id
            self.por_ruta[self._ruta(rutas[u.id])] = u.id
            nombres.setdefault(normalizar(u.nombre), []).append(u.id)
        # Solo se usa el nombre si es único (puede haber muchas «Balda 1»).
        self.por_nombre = {n: ids[0] for n, ids in nombres.items() if len(ids) == 1}

    @staticmethod
    def _ruta(valor: str) -> str:
        return "/".join(normalizar(p) for p in re.split(r"[›>/\\]", valor) if p.strip())

    def buscar(self, valor: str) -> int | None:
        clave = normalizar(valor)
        return self.por_codigo.get(clave) or self.por_ruta.get(self._ruta(valor)) or self.por_nombre.get(clave)


# ------------------------------------------------------------------ importación

def importar(con: sqlite3.Connection, tabla: Tabla, mapa: dict[int, str], tipo_por_defecto: int,
             ubicacion_por_defecto: int | None = None, omitir_repetidos: bool = True,
             progreso=None) -> Informe:
    """Crea un elemento por cada fila. ``mapa``: nº de columna -> destino (los de ``destinos``).

    Todo va en una transacción: si algo falla de forma inesperada no queda nada a medias.
    Las filas con problemas leves se importan igualmente y se anotan en ``Informe.avisos``.
    """
    if "titulo" not in mapa.values():
        raise ErrorImportar("Indica qué columna contiene el título.")
    lista_tipos = tipos.listar(con)
    por_id = {t.id: t for t in lista_tipos}
    if tipo_por_defecto not in por_id:
        raise ErrorImportar("El tipo elegido ya no existe.")
    tipo_por_nombre = {normalizar(t.nombre): t.id for t in lista_tipos}
    buscador = _Ubicaciones(con)
    estados = {normalizar(e): e for e in elementos.ESTADOS}
    informe = Informe()
    vistos: set[str] = set()  # identificadores ya importados en este mismo archivo
    columnas = sorted(mapa.items())

    with transaccion(con):
        for n, fila in enumerate(tabla.filas):
            numero = n + 2  # nº de fila tal y como se ve en Excel (la 1 son las cabeceras)
            if progreso and n % 200 == 0:
                progreso(n, len(tabla.filas))
            valores = {destino: [] for _, destino in columnas}
            for columna, destino in columnas:
                if destino and columna < len(fila) and fila[columna]:
                    valores[destino].append((tabla.cabeceras[columna], fila[columna]))

            def primero(destino: str) -> str:
                return valores.get(destino, [("", "")])[0][1] if valores.get(destino) else ""

            avisos: list[str] = []
            notas: list[str] = []

            # Tipo de la fila
            t = por_id[tipo_por_defecto]
            if primero("tipo"):
                tipo_id = tipo_por_nombre.get(normalizar(primero("tipo")))
                if tipo_id:
                    t = por_id[tipo_id]
                else:
                    avisos.append(f"tipo «{primero('tipo')}» desconocido: se usa «{t.nombre}»")

            titulo = " ".join(v for _, v in valores.get("titulo", [])).strip()
            if not titulo:
                informe.omitidos.append((numero, "sin título"))
                continue
            e = Elemento(tipo_id=t.id, titulo=titulo, ubicacion_id=ubicacion_por_defecto)
            e.subtitulo = primero("subtitulo")
            e.identificador = elementos.normalizar_identificador(primero("identificador"))
            if e.identificador:
                if omitir_repetidos and (e.identificador in vistos
                                         or elementos.buscar_por_identificador(con, e.identificador)):
                    informe.omitidos.append((numero, f"«{titulo}»: ya existe el identificador {e.identificador}"))
                    continue
                vistos.add(e.identificador)

            # Personas: el rol sale del nombre de la columna si coincide con uno del tipo
            # (columna «Traductor» -> rol Traductor); si no, el primer rol del tipo.
            roles = {normalizar(r): r for r in t.roles}
            rol_defecto = t.roles[0] if t.roles else "Autor"
            for cabecera, valor in valores.get("personas", []):
                e.personas += _personas(valor, roles.get(normalizar(cabecera), rol_defecto))

            if primero("anio"):
                anio = _anio(primero("anio"))
                if anio is not None:
                    e.anio = anio
                else:
                    notas.append(f"Año: {primero('anio')}")
            e.idioma = primero("idioma")
            if primero("estado"):
                estado = estados.get(normalizar(primero("estado")))
                if estado:
                    e.estado = estado
                else:
                    notas.append(f"Conservación: {primero('estado')}")
            if primero("valoracion"):
                valoracion = _valoracion(primero("valoracion"))
                if valoracion is not None:
                    e.valoracion = valoracion
                else:
                    notas.append(f"Valoración: {primero('valoracion')}")
            e.consumido = normalizar(primero("consumido")) in VALORES_SI
            for _, valor in valores.get("etiquetas", []):
                e.etiquetas += [x for x in re.split(r"[,;]", valor) if x.strip()]
            e.prestado_a = primero("prestado_a")
            if e.prestado_a:
                e.fecha_prestamo = date.today().isoformat()
            for destino, nombre in (("fecha_desde", "Desde"), ("fecha_hasta", "Hasta")):
                valor = primero(destino)
                if elementos.fecha_valida(valor):
                    setattr(e, destino, valor)
                else:
                    notas.append(f"{nombre}: {valor}")
            if e.fecha_desde and e.fecha_hasta and e.fecha_hasta < e.fecha_desde:
                notas.append(f"Periodo: {e.fecha_desde} – {e.fecha_hasta}")
                e.fecha_desde = e.fecha_hasta = ""
            e.lugar_evento = primero("lugar_evento")

            if primero("ubicacion"):
                ubicacion_id = buscador.buscar(primero("ubicacion"))
                if ubicacion_id:
                    e.ubicacion_id = ubicacion_id
                else:
                    avisos.append(f"ubicación «{primero('ubicacion')}» no encontrada")
                    notas.append(f"Ubicación indicada: {primero('ubicacion')}")

            # Campos propios del tipo; si este tipo no tiene ese campo, el dato va a las notas.
            campos_tipo = {normalizar(c.etiqueta): c for c in t.campos}
            campos_tipo.update({normalizar(c.clave): c for c in t.campos})
            for destino, lista in valores.items():
                if not destino.startswith(PREFIJO_CAMPO):
                    continue
                campo = campos_tipo.get(destino.removeprefix(PREFIJO_CAMPO))
                for cabecera, valor in lista:
                    if campo and campo.id not in e.valores:
                        e.valores[campo.id] = valor
                    else:
                        notas.append(f"{cabecera}: {valor}")

            # Notas: la primera columna tal cual; las demás con su nombre delante.
            lista_notas = valores.get("notas", [])
            partes = [lista_notas[0][1]] if lista_notas else []
            partes += [f"{c}: {v}" for c, v in lista_notas[1:]]
            e.notas = "\n".join(partes + notas)

            try:
                informe.creados.append(elementos.guardar(con, e))
            except ErrorElemento as error:
                informe.omitidos.append((numero, f"«{titulo}»: {error}"))
                continue
            if avisos:
                informe.avisos.append((numero, f"«{titulo}»: " + "; ".join(avisos)))
    if progreso:
        progreso(len(tabla.filas), len(tabla.filas))
    return informe


# ------------------------------------------------------------------ plantilla

COLUMNAS_PLANTILLA = ["Tipo", "Título", "Autor", "Año", "ISBN", "Editorial", "Ubicación", "Etiquetas", "Notas"]
EJEMPLO_PLANTILLA = ["Libro", "Cien años de soledad", "Gabriel García Márquez", "1967", "978-84-376-0494-7",
                     "Cátedra", "PB", "novela, clásicos", "Ejemplo: borra esta fila"]


def escribir_plantilla(destino: Path | str) -> None:
    """Hoja de Excel vacía con las columnas habituales y una hoja de instrucciones."""
    import openpyxl
    from openpyxl.styles import Font

    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Elementos"
    hoja.append(COLUMNAS_PLANTILLA)
    hoja.append(EJEMPLO_PLANTILLA)
    for celda in hoja[1]:
        celda.font = Font(bold=True)
    for letra, ancho in zip("ABCDEFGHI", (12, 36, 28, 8, 20, 18, 22, 22, 30)):
        hoja.column_dimensions[letra].width = ancho
    hoja.freeze_panes = "A2"
    ayuda = libro.create_sheet("Instrucciones")
    for linea in (
        "Cómo rellenar esta hoja para importarla en LibriDomus",
        "",
        "• Una fila por objeto físico. Solo el Título es obligatorio.",
        "• Tipo: Libro, Disco, Película, Revista / Cómic... (si se deja vacío se usa el que elijas al importar).",
        "• Autor: varias personas separadas por punto y coma. Para indicar el papel: «Luis Gil (Traductor)».",
        "• Ubicación: el código de la etiqueta (p. ej. PB-SAL-EA-B3) o la ruta (Planta baja > Salón > Estantería A).",
        "• Etiquetas: separadas por comas.",
        "• Puedes añadir más columnas (Idioma, Conservación, Notas, Nº de páginas...): al importar eliges a qué "
        "dato va cada una. Lo que no encaje se guarda en las notas.",
        "• La fila 2 es un ejemplo: bórrala antes de importar.",
        "• LibriDomus solo lee la primera hoja (Elementos).",
    ):
        ayuda.append([linea])
    ayuda["A1"].font = Font(bold=True, size=13)
    ayuda.column_dimensions["A"].width = 110
    libro.save(str(destino))
